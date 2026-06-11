#!/usr/bin/env python3
"""
Generate the RenovateAI audit deliverables:
  - docs/deliverables/RenovateAI_Audit_Report.pdf
  - docs/deliverables/RenovateAI_Audit_Report.docx
  - docs/deliverables/RenovateAI_Feature_Tracker.xlsx

Run with: python3 scripts/generate_reports.py
"""
import os
from datetime import date

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, ListFlowable, ListItem
)

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "docs", "deliverables")
os.makedirs(OUT_DIR, exist_ok=True)
TODAY = date.today().isoformat()

# ---------------------------------------------------------------------------
# Shared data
# ---------------------------------------------------------------------------

FEATURES = [
    (1, "Multi-tenant foundation", "Each contractor gets an isolated workspace; data never crosses tenants", "Done", "1"),
    (2, "Auth (Clerk)", "Sign up / sign in, session management", "Done", "1"),
    (3, "Onboarding wizard", "New contractor picks a slug -> workspace auto-provisioned", "Done", "1"),
    (4, "BrandKit", "Logo, colors, brand voice, bio, license/insurance shown in widget", "Done", "2"),
    (5, "Homeowner widget", "Embeddable page where homeowners upload photos & chat", "Done", "2"),
    (6, "AI photo renders (FLUX/Fal.ai)", "Upload kitchen photo -> AI shows renovated version in ~60s", "Done", "2"),
    (7, "AI chat assistant (Claude)", "24/7 assistant answering using contractor's own info", "Done", "3"),
    (8, "Real dashboard stats", "Live lead/render counts on dashboard home", "Done", "3"),
    (9, "Embed snippet generator", "Copy-paste iframe code for contractor's site", "Done", "3"),
    (10, "Tenant provisioning API", "Automated workspace creation on signup", "Done", "4"),
    (11, "Stripe billing", "Checkout, customer portal, plan upgrades via webhook", "Done", "4"),
    (12, "Follow-up emails", "Automatic reminder emails to leads who go quiet", "Done", "4"),
    (13, "Security hardening pass #1", "Rate limiting, validation, error boundaries, admin panel, theming", "Done", "5"),
    (14, "Plan limits + usage metering", "Starter/Growth/Pro render & lead caps enforced server-side", "Done", "6"),
    (15, "GoHighLevel CRM sync", "Leads pushed to contractor's GHL pipeline", "Done", "6"),
    (16, "AI memory / RAG (pgvector)", "Contractor uploads docs/FAQs; chatbot retrieves relevant chunks", "Done", "6.5"),
    (17, "SEO module", "AI-generated local landing pages with JSON-LD schema", "Done", "6.5"),
    (18, "GHL OAuth connect flow", "One-click CRM connection (no manual API keys)", "Done", "7"),
    (19, "Trial reminder emails", "\"Your trial ends in N days\" automated emails", "Done", "7"),
    (20, "Chat trimming + usage banner", "Keeps token costs bounded; shows usage in dashboard", "Done", "7"),
    (21, "AI Scope Analyzer", "Turns chat into structured project spec", "Done", "8"),
    (22, "AI Budget Estimator", "Price range using contractor's own pricing rules", "Done", "8"),
    (23, "AI Proposal Generator", "One-click web proposal, e-signature, Stripe deposit", "Done", "9"),
    (24, "Review & Testimonial Assistant", "4-5 star -> public review links; 1-3 star -> private alert + AI draft reply", "Done", "10"),
    (25, "Tamper-evident audit log", "Every signature/payment/settings/integration change logged with IP + actor", "Done", "10.5"),
    (26, "Email XSS hardening", "User text escaped before going into HTML emails", "Done", "10.5"),
    (27, "Public API rate limiting", "/api/v1/leads and /api/v1/renders capped per-IP", "Done", "10.5"),
    (28, "Public API key auth", "/api/v1/* requires per-tenant Bearer key, managed in Settings", "Done", "10.6"),
    (29, "Render lookup tenant scoping", "GET /api/v1/renders?id= requires owning tenant's API key", "Done", "10.6"),
]

VERIFICATION = [
    ("npm run typecheck (tsc --noEmit)", "PASS", "0 errors"),
    ("npm run lint (eslint)", "PASS", "0 errors, 0 warnings"),
    ("npm run build (next build, production)", "PASS", "All ~45 routes compiled successfully"),
    ("drizzle-kit generate (schema -> migration)", "PASS", "Migration 0001_shallow_vivisector.sql generated"),
]

ROADMAP = [
    ("11", "Multi-CRM hub", "HubSpot + Jobber adapters, field-mapping UI, sync dead-letter view", "Not started"),
    ("12", "Lead intelligence", "AI lead scoring (0-100), hot-lead SMS routing, scoring feedback loop", "Not started"),
    ("13", "Launch polish", "sitemap.xml/robots.txt, Playwright funnel test, public launch checklist", "Not started"),
    ("M4", "SAM2 region-select rendering", "Tap-to-select regions, ControlNet depth/edge conditioning", "Not started"),
    ("M5", "Voice + speed-to-lead", "Vapi/Bland AI voice agent, SMS follow-up channel", "Not started"),
    ("M6", "Agency / white-label tier", "Sub-accounts, white-label theming, GHL Snapshot listing", "Not started"),
    ("M7", "Vertical: bathroom remodeling", "New vertical_config, style references, prompt templates", "Not started"),
    ("M8", "Custom LoRA + ROI dashboard", "Per-tenant style training, revenue-impact dashboard", "Not started"),
    ("M9", "Vertical: landscaping", "Same pattern as bathroom, new style set", "Not started"),
]

SECURITY_ITEMS = [
    ("Tenant isolation", "Every query scoped by tenantId via requireTenant()", "Implemented"),
    ("Input validation", "Zod schemas on every API route", "Implemented"),
    ("Stripe webhook verification", "stripe.webhooks.constructEvent signature check", "Implemented"),
    ("CRM credential encryption", "AES-256-GCM at rest (ENCRYPTION_KEY)", "Implemented"),
    ("Photo privacy", "EXIF stripped on upload", "Implemented"),
    ("Rate limiting", "Per-IP limits on all public endpoints", "Implemented"),
    ("Unguessable tokens", "Proposal & review links use random access tokens", "Implemented"),
    ("SEO HTML safety", "AI output rendered via escaped structured templates, never raw HTML", "Implemented"),
    ("Audit logging", "audit_log table: signatures, payments, settings, integrations, API keys", "Implemented"),
    ("Email XSS protection", "escapeHtml() on all user-controlled email fields", "Implemented"),
    ("Public API authentication", "/api/v1/* requires hashed per-tenant Bearer key", "Implemented"),
    ("Render lookup scoping", "GET /api/v1/renders?id= scoped to API-key tenant", "Implemented"),
    ("Distributed rate limiting", "Move in-process limiter to Upstash Redis for multi-region", "Recommended"),
    ("Structured log export", "Ship audit_log to Axiom/BetterStack", "Recommended"),
    ("CSP headers on widget", "Content-Security-Policy for embed surface", "Recommended"),
    ("Penetration test", "Before Agency/white-label tier launch", "Recommended"),
]

MONETIZATION = [
    ("Tiered metering", "Starter/Growth/Pro/Agency caps on renders & leads", "Already enforced"),
    ("Annual plans", "2 months free for annual commitment", "Improves cashflow & churn"),
    ("Setup fee", "$497-$997 one-time white-glove onboarding", "Pure margin"),
    ("Add-ons", "SEO page packs, custom LoRA + maintenance, AI voice agent", "ARPU expansion"),
    ("Agency / white-label tier", "1 sale = 10+ contractor sub-accounts", "Highest leverage channel"),
    ("Referral program", "30% rev-share, contractor refers contractor", "Low-CAC growth"),
    ("API access add-on", "Now authenticated -- sell as Pro+/Agency perk (Zapier, etc.)", "New, enabled this session"),
]

# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------

def build_docx():
    doc = Document()

    title = doc.add_heading("RenovateAI — Product Audit & Status Report", level=0)
    sub = doc.add_paragraph(f"Generated {TODAY}  ·  Branch claude/funny-lamport-MZCwj")
    sub.runs[0].italic = True

    doc.add_heading("1. Feature Completion (29/29 shipped)", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["#", "Feature", "Description", "Week"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for num, name, descr, status, week in FEATURES:
        row = table.add_row().cells
        row[0].text = str(num)
        row[1].text = f"{name}  [DONE]"
        row[2].text = descr
        row[3].text = week

    doc.add_heading("2. Verification Evidence", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Check", "Result", "Notes"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for check, result, notes in VERIFICATION:
        row = table.add_row().cells
        row[0].text = check
        row[1].text = result
        row[2].text = notes

    doc.add_heading("3. Security Posture", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Control", "Detail", "Status"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for ctrl, detail, status in SECURITY_ITEMS:
        row = table.add_row().cells
        row[0].text = ctrl
        row[1].text = detail
        row[2].text = status

    doc.add_heading("4. Roadmap — Weeks 11+ and Advanced Level", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Week/Month", "Theme", "Deliverables", "Status"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for wk, theme, deliv, status in ROADMAP:
        row = table.add_row().cells
        row[0].text = wk
        row[1].text = theme
        row[2].text = deliv
        row[3].text = status

    doc.add_heading("5. Monetization Levers", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Lever", "Detail", "Why it matters"]):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
    for lever, detail, why in MONETIZATION:
        row = table.add_row().cells
        row[0].text = lever
        row[1].text = detail
        row[2].text = why

    doc.add_heading("6. Architecture Summary", level=1)
    doc.add_paragraph(
        "Single Next.js 16 (App Router, React 19, TypeScript strict) application "
        "containing both frontend and backend:"
    )
    bullets = [
        "/dashboard/* — contractor admin (leads, renders, proposals, reviews, SEO, settings, admin)",
        "/embed/[tenantSlug] — homeowner-facing widget, white-labeled per contractor",
        "/p/[token], /r/[token], /s/[tenantSlug]/[slug] — public proposal, review, and SEO pages",
        "~32 API routes under /api/*, Drizzle ORM over Postgres + pgvector",
        "Inngest for async jobs: renders, follow-up emails, CRM sync",
        "Deploy target: Vercel (app) + Neon (DB) + Cloudflare R2 (images)",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    doc.add_heading("7. Bottom Line", level=1)
    doc.add_paragraph(
        "All 10 weeks of the original MVP plan plus a 3-stage hardening pass are complete, "
        "verified via typecheck/lint/production build, and pushed to GitHub. The technical "
        "foundation for speed, SEO, and AI-search visibility is solid (8-8.5/10) — the next "
        "gains come from content volume, backlinks, and the Week 11-13 roadmap items above."
    )

    path = os.path.join(OUT_DIR, "RenovateAI_Audit_Report.docx")
    doc.save(path)
    return path

# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------

def build_xlsx():
    wb = Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    done_fill = PatternFill("solid", fgColor="DCFCE7")
    todo_fill = PatternFill("solid", fgColor="FEF9C3")
    warn_fill = PatternFill("solid", fgColor="FEE2E2")
    thin = Side(border_style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border

    # --- Sheet 1: Features ---
    ws = wb.active
    ws.title = "Features"
    ws.append(["#", "Feature", "Description", "Status", "Week"])
    style_header(ws, 5)
    for num, name, descr, status, week in FEATURES:
        ws.append([num, name, descr, status, week])
        r = ws.max_row
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = border
            ws.cell(row=r, column=c).alignment = wrap
        ws.cell(row=r, column=4).fill = done_fill
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 8
    ws.freeze_panes = "A2"

    # --- Sheet 2: Verification / Test Proof ---
    ws2 = wb.create_sheet("Test Proof")
    ws2.append(["Check", "Result", "Notes"])
    style_header(ws2, 3)
    for check, result, notes in VERIFICATION:
        ws2.append([check, result, notes])
        r = ws2.max_row
        for c in range(1, 4):
            ws2.cell(row=r, column=c).border = border
            ws2.cell(row=r, column=c).alignment = wrap
        ws2.cell(row=r, column=2).fill = done_fill
    ws2.append([])
    ws2.append(["Manual end-to-end funnel checklist (see docs/PROJECT_STATUS.md section 2.3)"])
    funnel_steps = [
        "Sign up -> onboarding -> workspace created",
        "BrandKit configured (logo, colors, voice)",
        "Embed widget loads at /embed/<slug>",
        "Photo upload + AI render completes",
        "Chat assistant responds using BrandKit voice + RAG",
        "Lead captured in dashboard",
        "Follow-up email sent (Resend)",
        "Scope -> Estimate -> Proposal -> Send",
        "Proposal signed (e-sign) + Stripe deposit paid",
        "Review request: 5-star -> public links, 2-star -> private owner alert",
        "API key created in Settings; /api/v1/leads call with Bearer key returns 201",
        "/api/v1/leads call WITHOUT Bearer key returns 401 (auth enforced)",
    ]
    for step in funnel_steps:
        ws2.append(["", "To run on a configured environment", step])
        r = ws2.max_row
        for c in range(1, 4):
            ws2.cell(row=r, column=c).border = border
            ws2.cell(row=r, column=c).alignment = wrap
        ws2.cell(row=r, column=2).fill = todo_fill
    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 28
    ws2.column_dimensions["C"].width = 60

    # --- Sheet 3: Security ---
    ws3 = wb.create_sheet("Security")
    ws3.append(["Control", "Detail", "Status"])
    style_header(ws3, 3)
    for ctrl, detail, status in SECURITY_ITEMS:
        ws3.append([ctrl, detail, status])
        r = ws3.max_row
        for c in range(1, 4):
            ws3.cell(row=r, column=c).border = border
            ws3.cell(row=r, column=c).alignment = wrap
        ws3.cell(row=r, column=3).fill = done_fill if status == "Implemented" else warn_fill
    ws3.column_dimensions["A"].width = 28
    ws3.column_dimensions["B"].width = 65
    ws3.column_dimensions["C"].width = 14

    # --- Sheet 4: Roadmap ---
    ws4 = wb.create_sheet("Roadmap")
    ws4.append(["Week/Month", "Theme", "Deliverables", "Status"])
    style_header(ws4, 4)
    for wk, theme, deliv, status in ROADMAP:
        ws4.append([wk, theme, deliv, status])
        r = ws4.max_row
        for c in range(1, 5):
            ws4.cell(row=r, column=c).border = border
            ws4.cell(row=r, column=c).alignment = wrap
        ws4.cell(row=r, column=4).fill = todo_fill
    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 30
    ws4.column_dimensions["C"].width = 65
    ws4.column_dimensions["D"].width = 14

    # --- Sheet 5: Monetization ---
    ws5 = wb.create_sheet("Monetization")
    ws5.append(["Lever", "Detail", "Why it matters"])
    style_header(ws5, 3)
    for lever, detail, why in MONETIZATION:
        ws5.append([lever, detail, why])
        r = ws5.max_row
        for c in range(1, 4):
            ws5.cell(row=r, column=c).border = border
            ws5.cell(row=r, column=c).alignment = wrap
    ws5.column_dimensions["A"].width = 26
    ws5.column_dimensions["B"].width = 55
    ws5.column_dimensions["C"].width = 30

    path = os.path.join(OUT_DIR, "RenovateAI_Feature_Tracker.xlsx")
    wb.save(path)
    return path

# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def build_pdf():
    path = os.path.join(OUT_DIR, "RenovateAI_Audit_Report.pdf")
    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=1.5*cm, bottomMargin=1.5*cm,
                             leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], fontSize=20, spaceAfter=4)
    h1 = ParagraphStyle("H1X", parent=styles["Heading1"], fontSize=14, spaceBefore=14, spaceAfter=6,
                         textColor=colors.HexColor("#1F2937"))
    body = ParagraphStyle("BodyX", parent=styles["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("SmallX", parent=styles["BodyText"], fontSize=8, leading=10)

    story = []
    story.append(Paragraph("RenovateAI — Product Audit &amp; Status Report", title_style))
    story.append(Paragraph(f"Generated {TODAY} · Branch claude/funny-lamport-MZCwj", styles["Italic"]))
    story.append(Spacer(1, 10))

    story.append(Paragraph("1. Feature Completion — 29 / 29 shipped", h1))
    data = [["#", "Feature", "Description", "Wk"]]
    for num, name, descr, status, week in FEATURES:
        data.append([str(num), Paragraph(name, small), Paragraph(descr, small), week])
    t = Table(data, colWidths=[1*cm, 4.5*cm, 10.5*cm, 1*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0FDF4")]),
    ]))
    story.append(t)
    story.append(PageBreak())

    story.append(Paragraph("2. Verification Evidence", h1))
    data = [["Check", "Result", "Notes"]]
    for check, result, notes in VERIFICATION:
        data.append([Paragraph(check, small), result, Paragraph(notes, small)])
    t = Table(data, colWidths=[7*cm, 2*cm, 8*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)

    story.append(Paragraph("3. Security Posture", h1))
    data = [["Control", "Detail", "Status"]]
    for ctrl, detail, status in SECURITY_ITEMS:
        data.append([Paragraph(ctrl, small), Paragraph(detail, small), status])
    t = Table(data, colWidths=[4.5*cm, 10.5*cm, 2*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)
    story.append(PageBreak())

    story.append(Paragraph("4. Roadmap — Weeks 11+ and Advanced Level", h1))
    data = [["Wk/Mo", "Theme", "Deliverables", "Status"]]
    for wk, theme, deliv, status in ROADMAP:
        data.append([wk, Paragraph(theme, small), Paragraph(deliv, small), status])
    t = Table(data, colWidths=[1.5*cm, 3.5*cm, 9.5*cm, 2.5*cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)

    story.append(Paragraph("5. Monetization Levers", h1))
    items = [Paragraph(f"<b>{lever}</b> — {detail} <i>({why})</i>", body) for lever, detail, why in MONETIZATION]
    story.append(ListFlowable([ListItem(i) for i in items], bulletType="bullet"))

    story.append(Paragraph("6. Bottom Line", h1))
    story.append(Paragraph(
        "All 10 weeks of the original MVP plan plus a 3-stage hardening pass are complete, "
        "verified via typecheck/lint/production build, and pushed to GitHub. The technical "
        "foundation for speed, SEO, and AI-search visibility is solid (8-8.5/10) — the next "
        "gains come from content volume, backlinks, and the Week 11-13 roadmap items above.",
        body))

    doc.build(story)
    return path

if __name__ == "__main__":
    p1 = build_docx()
    p2 = build_xlsx()
    p3 = build_pdf()
    for p in (p1, p2, p3):
        print("Wrote", p)
